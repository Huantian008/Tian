from copy import deepcopy
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(r"C:\Users\y2003\Downloads\03+张三+202499770310（模版）.docx")
OUT = ROOT / "deliverables" / "11+杨文天+202599770425.docx"
ZIP_OUT = ROOT / "deliverables" / "11+杨文天+202599770425.zip"
SCREENSHOTS = ROOT / "docs" / "screenshots"


def set_run_font(run, size=11, bold=False, color=None, font="宋体"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def replace_paragraph_text(paragraph, text):
    if not paragraph.runs:
        paragraph.add_run("")
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def remove_body_after_paragraph(doc, keep_index):
    marker = doc.paragraphs[keep_index]._element
    body = doc._body._element
    remove = False
    for child in list(body):
        if child is marker:
            remove = True
            continue
        if not remove:
            continue
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)


def add_para(doc, text="", size=11, bold=False, align=None, color=None, style=None):
    p = doc.add_paragraph(style=style or "Normal")
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    return p


def add_heading(doc, text, level=2, center=False):
    style = "Heading 2" if level <= 2 else "Heading 3"
    p = doc.add_paragraph(style=style)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=16 if center else 12, bold=True, font="黑体")
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    return p


def set_cell_text(cell, text, bold=False, shade=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(str(text)) <= 16 else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    set_run_font(run, size=10, bold=bold)
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_after = Pt(0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "left", "bottom", "right"):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), "120")
        node.set(qn("w:type"), "dxa")
    if shade:
        shd = tc_pr.find(qn("w:shd"))
        if shd is None:
            shd = OxmlElement("w:shd")
            tc_pr.append(shd)
        shd.set(qn("w:fill"), shade)


def set_table_widths(table, widths_cm):
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            row.cells[idx].width = Cm(width)
            tc_pr = row.cells[idx]._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(int(width * 567)))


def add_table(doc, headers, rows, widths_cm=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, shade="E8EEF5")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    if widths_cm:
        set_table_widths(table, widths_cm)
    doc.add_paragraph()
    return table


def add_code(doc, title, path, start_marker=None, max_lines=24):
    add_para(doc, title, bold=True, style="论文内容")
    text = (ROOT / path).read_text(encoding="utf-8")
    if start_marker and start_marker in text:
        text = text[text.index(start_marker) :]
    lines = text.strip().splitlines()[:max_lines]
    p = doc.add_paragraph(style="HTML Preformatted")
    for line in lines:
        run = p.add_run(line[:96] + "\n")
        set_run_font(run, size=8, font="Consolas")
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(8)


def add_screenshot(doc, title, file_name):
    p = add_para(doc, title, align=WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.keep_with_next = True
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = pic.add_run()
    run.add_picture(str(SCREENSHOTS / file_name), width=Cm(14.8))
    pic.paragraph_format.space_after = Pt(8)


def clone_template_table_style(doc):
    if not doc.tables:
        return None
    template_tbl = deepcopy(doc.tables[0]._tbl)
    return template_tbl


def build_body(doc):
    add_heading(doc, "第1章 系统设计", center=True)
    add_heading(doc, "1.1  系统总体设计")
    add_para(
        doc,
        "学生管理系统采用前后端分离结构。前端使用 Vue 3、Vite 和 Element Plus 构建后台管理界面，后端使用 Spring Boot 3 和 MyBatis 提供 REST API，数据库使用 MySQL 保存管理员、学生、课程和成绩选课数据。系统主要包括用户登录、首页统计、学生信息管理、课程信息管理、成绩与选课管理等模块。",
    )
    add_table(
        doc,
        ["层次", "技术/模块", "主要职责"],
        [
            ["表现层", "Vue 3 + Element Plus", "登录、首页统计、学生/课程/成绩管理页面"],
            ["接口层", "Spring Boot Controller", "接收前端请求，返回统一 JSON 数据"],
            ["业务/持久层", "Service 思路 + MyBatis Mapper", "完成数据校验、查询、新增、修改和删除"],
            ["数据层", "MySQL", "保存管理员、学生、课程、成绩与选课记录"],
        ],
        [3.0, 5.0, 7.0],
    )
    add_para(doc, "图1 系统功能结构图", align=WD_ALIGN_PARAGRAPH.CENTER)

    add_heading(doc, "1.2   数据库表设计", level=3)
    add_para(
        doc,
        "数据库名为 student_management_system，包含 admin_user、student、course、score 四张核心表。admin_user 存储管理员登录信息；student 存储学生基础信息；course 存储课程信息；score 通过学生 ID 和课程 ID 建立选课与成绩关联。",
        style="论文内容",
    )
    add_para(doc, "表1  管理员用户表结构", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_table(
        doc,
        ["字段", "数据类型", "字段名", "备注"],
        [
            ["id", "bigint", "主键ID", "管理员主键"],
            ["username", "varchar(50)", "用户名", "登录账号"],
            ["password", "varchar(100)", "密码", "演示环境使用明文密码"],
            ["real_name", "varchar(50)", "真实姓名", "管理员姓名"],
            ["role", "varchar(30)", "角色", "系统角色"],
        ],
        [3.2, 3.6, 3.5, 4.5],
    )
    add_para(doc, "表2  学生信息表结构", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_table(
        doc,
        ["字段", "数据类型", "字段名", "备注"],
        [
            ["id", "bigint", "主键ID", "学生主键"],
            ["student_no", "varchar(30)", "学号", "唯一"],
            ["name", "varchar(50)", "姓名", "学生姓名"],
            ["gender", "varchar(10)", "性别", "男/女"],
            ["age", "int", "年龄", "学生年龄"],
            ["class_name", "varchar(80)", "班级", "所在班级"],
            ["major", "varchar(100)", "专业", "所属专业"],
            ["phone/email", "varchar", "联系方式", "电话和邮箱"],
        ],
        [3.2, 3.6, 3.5, 4.5],
    )
    add_para(doc, "表3  课程信息表结构", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_table(
        doc,
        ["字段", "数据类型", "字段名", "备注"],
        [
            ["id", "bigint", "主键ID", "课程主键"],
            ["course_no", "varchar(30)", "课程编号", "唯一"],
            ["course_name", "varchar(100)", "课程名称", "课程中文名称"],
            ["credit", "decimal(4,1)", "学分", "课程学分"],
            ["teacher", "varchar(50)", "任课教师", "课程教师"],
            ["semester", "varchar(30)", "学期", "开课学期"],
        ],
        [3.2, 3.6, 3.5, 4.5],
    )
    add_para(doc, "表4  成绩与选课表结构", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_table(
        doc,
        ["字段", "数据类型", "字段名", "备注"],
        [
            ["id", "bigint", "主键ID", "成绩主键"],
            ["student_id", "bigint", "学生ID", "关联学生表"],
            ["course_id", "bigint", "课程ID", "关联课程表"],
            ["score", "decimal(5,1)", "成绩", "课程成绩"],
            ["semester", "varchar(30)", "学期", "成绩所属学期"],
        ],
        [3.2, 3.6, 3.5, 4.5],
    )

    add_heading(doc, "第2章 系统开发与实现", center=True)
    add_heading(doc, "2.1  系统技术选型")
    add_para(
        doc,
        "后端选择 Spring Boot 3 作为 Java Web 框架，用注解和自动配置减少传统 SSM 项目的 XML 配置量；持久层使用 MyBatis，便于编写 SQL 并完成对象映射；数据库使用 MySQL，满足关系型数据存储要求；前端使用 Vue 3、Vite 和 Element Plus，可以快速完成后台管理页面和表单交互。",
    )
    add_heading(doc, "2.2  开发环境建立")
    add_para(
        doc,
        "系统在 Windows 环境下开发和测试，后端使用 Java 21，数据库使用本机 MySQL，前端通过 Node.js 和 Vite 启动开发服务器。项目内置 Maven Wrapper，可不依赖本机全局 Maven 直接运行后端。",
    )
    add_table(
        doc,
        ["类别", "环境"],
        [
            ["操作系统", "Windows"],
            ["后端语言", "Java 21"],
            ["后端框架", "Spring Boot 3、MyBatis"],
            ["数据库", "MySQL 8.0.45，账号 root，密码 ywt123YWT"],
            ["前端框架", "Vue 3、Vite、Element Plus、Axios"],
            ["开发工具", "IntelliJ IDEA / VS Code / Codex"],
        ],
        [4.0, 11.0],
    )
    add_heading(doc, "2.3  主要模块的实现")
    add_heading(doc, "2.3.1  用户登录实现界面", level=3)
    add_para(doc, "管理员输入账号和密码后，前端调用 /api/auth/login 接口。后端查询 admin_user 表并校验密码，校验通过后返回用户信息和简单 token。")
    add_screenshot(doc, "图2 用户登录实现界面", "01-login.png")
    add_code(doc, "其核心代码如下：", "backend/src/main/java/com/ywt/sms/controller/AuthController.java", "public ApiResponse<LoginResponse> login")

    add_heading(doc, "2.3.2  首页统计实现界面", level=3)
    add_para(doc, "首页统计展示学生总数、课程总数、选课记录数和平均成绩，便于管理员进入系统后快速掌握当前数据概况。")
    add_screenshot(doc, "图3 首页统计实现界面", "02-dashboard.png")
    add_code(doc, "其核心代码如下：", "backend/src/main/java/com/ywt/sms/controller/DashboardController.java", "public ApiResponse<DashboardSummary> summary")

    add_heading(doc, "2.3.3  学生信息管理界面", level=3)
    add_para(doc, "学生管理模块提供按关键字查询、新增、编辑和删除功能，维护学号、姓名、性别、年龄、班级、专业、电话和邮箱等基础信息。")
    add_screenshot(doc, "图4 学生信息管理界面", "03-students.png")
    add_code(doc, "其核心代码如下：", "backend/src/main/java/com/ywt/sms/mapper/StudentMapper.java", "@Select")

    add_heading(doc, "2.3.4  课程与成绩管理界面", level=3)
    add_para(doc, "课程管理维护课程编号、课程名称、学分、任课教师和学期；成绩管理维护学生与课程的选课关系，并支持录入、修改和删除课程成绩。")
    add_screenshot(doc, "图5 课程管理界面", "04-courses.png")
    add_screenshot(doc, "图6 成绩管理界面", "05-scores.png")
    add_code(doc, "其核心代码如下：", "frontend/src/api.js", "export const api")

    add_heading(doc, "第3章 课程总结", center=True)
    add_para(
        doc,
        "本次 WEB 框架技术期末作品以学生管理系统为主题，完成了从数据库设计、后端接口开发到前端页面展示的完整流程。项目采用 Spring Boot、MyBatis、MySQL、Vue 3 和 Element Plus 等技术，实现了管理员登录、首页统计、学生管理、课程管理、成绩管理等功能。通过这个项目，我进一步熟悉了前后端分离项目的开发方式，理解了 REST API、数据库表关联、Axios 请求封装和组件化页面开发之间的配合关系。",
    )
    add_para(
        doc,
        "在开发过程中，重点解决了数据库初始化、接口联调、跨域访问、表单校验、列表分页和页面截图整理等问题。后端通过 Maven Wrapper 运行，降低了环境依赖；前端通过 Vite 启动，提高了开发调试效率。数据库脚本可以重复初始化演示数据，便于系统展示和验收。",
    )
    add_para(
        doc,
        "系统仍有可以继续完善的地方，例如登录密码可改为加密存储，接口可增加更严格的权限校验，成绩统计可增加图表分析，前端页面也可以继续优化响应式布局。总体来看，本作品完成了课程要求中的主要功能，并将课堂中的 Java Web、数据库和前端框架知识整合到了一个可运行的管理系统中。",
    )


def update_cover(doc):
    replacements = {
        16: "学    院：         信息工程学院           ",
        17: "专    业：        计算机科学与技术        ",
        18: "班    级：        23专升本计算机7班        ",
        19: "学    号：  202599770425      姓   名：  杨文天      ",
        20: "任课教师1：           肖磊              ",
        23: "作品名称：学生管理系统",
    }
    for idx, text in replacements.items():
        replace_paragraph_text(doc.paragraphs[idx], text)


def build_zip():
    excluded_parts = {"node_modules", "target", "dist", "package-staging", ".git", ".idea"}
    with ZipFile(ZIP_OUT, "w", ZIP_DEFLATED) as zf:
        for path in ROOT.rglob("*"):
            if path.is_dir():
                continue
            rel = path.relative_to(ROOT)
            if any(part in excluded_parts or part.startswith("report-render") for part in rel.parts):
                continue
            if path == ZIP_OUT:
                continue
            zf.write(path, rel.as_posix())


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(str(TEMPLATE))
    update_cover(doc)
    remove_body_after_paragraph(doc, 23)
    build_body(doc)
    doc.save(OUT)
    build_zip()
    print(OUT)
    print(ZIP_OUT)


if __name__ == "__main__":
    build()
